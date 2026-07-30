from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "MANUAL_USUARIO.md"
OUTPUT = ROOT / "Manual_de_usuario_Fichajes.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
MUTED = RGBColor(90, 98, 110)
ACCENT = RGBColor(167, 191, 38)


def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_free_numbering(document):
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    next_abstract = max(abstract_ids, default=0) + 1
    next_num = max(num_ids, default=0) + 1

    def add_definition(num_format, level_text, font=None):
        nonlocal next_abstract, next_num
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(next_abstract))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)

        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_format)
        level.append(fmt)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), level_text)
        level.append(text)
        justification = OxmlElement("w:lvlJc")
        justification.set(qn("w:val"), "left")
        level.append(justification)

        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ppr.append(tabs)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "540")
        indent.set(qn("w:hanging"), "270")
        ppr.append(indent)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.append(spacing)
        level.append(ppr)

        if font:
            rpr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            rpr.append(fonts)
            level.append(rpr)

        abstract.append(level)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(next_num))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(next_abstract))
        num.append(abstract_ref)
        numbering.append(num)

        result = next_num
        next_abstract += 1
        next_num += 1
        return result

    return add_definition("bullet", "•", "Calibri"), add_definition("decimal", "%1.")


def new_number_instance(document, abstract_num_id):
    numbering = document.part.numbering_part.element
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    next_num = max(num_ids, default=0) + 1
    source_num = next(
        node for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == abstract_num_id
    )
    abstract_id = source_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(next_num))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract_id)
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return next_num


def apply_numbering(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(level)
    num_pr.append(num)
    ppr.append(num_pr)


def add_inline(paragraph, text):
    for part in re.split(r"(`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Consolas", 9.5, NAVY)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E8EEF5")
            run._element.get_or_add_rPr().append(shading)
        else:
            run = paragraph.add_run(part)
            set_font(run)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


def set_page_furniture(section, unlink=False):
    if unlink:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        section.even_page_header.is_linked_to_previous = False
        section.even_page_footer.is_linked_to_previous = False

    for header_part in (section.header, section.even_page_header):
        header = header_part.paragraphs[0]
        header._p.clear_content()
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(header.add_run("FICHAJES  |  MANUAL DE USUARIO"), size=8.5, color=MUTED, bold=True)

    for footer_part in (section.footer, section.even_page_footer):
        footer = footer_part.paragraphs[0]
        footer._p.clear_content()
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(footer.add_run("Página "), size=8.5, color=MUTED)
        add_page_field(footer)


document = Document()
document.settings.odd_and_even_pages_header_footer = True
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = document.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
]:
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

set_page_furniture(section)

document.core_properties.title = "Manual de usuario de Fichajes"
document.core_properties.subject = "Guía de uso de Fichajes v2.15.0"
document.core_properties.author = "Fichajes"

document.add_paragraph().paragraph_format.space_after = Pt(92)
kicker = document.add_paragraph()
kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
kicker.paragraph_format.space_after = Pt(16)
set_font(kicker.add_run("GUÍA DE REFERENCIA"), size=10, color=ACCENT, bold=True)

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(10)
set_font(title.add_run("Manual de usuario"), size=30, color=NAVY, bold=True)

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(58)
set_font(subtitle.add_run("Fichajes"), size=18, color=BLUE, bold=True)

metadata = document.add_paragraph()
metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
metadata.paragraph_format.space_after = Pt(4)
set_font(metadata.add_run("Versión de la aplicación 2.15.0"), size=11, color=MUTED, bold=True)
date_line = document.add_paragraph()
date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(date_line.add_run("Actualizado el 30 de julio de 2026"), size=10, color=MUTED)

contents_section = document.add_section(WD_SECTION.NEW_PAGE)
contents_section.page_width = Inches(8.5)
contents_section.page_height = Inches(11)
contents_section.top_margin = Inches(1)
contents_section.right_margin = Inches(1)
contents_section.bottom_margin = Inches(1)
contents_section.left_margin = Inches(1)
contents_section.header_distance = Inches(0.492)
contents_section.footer_distance = Inches(0.492)
set_page_furniture(contents_section, unlink=True)
contents_spacer = document.add_paragraph()
contents_spacer.paragraph_format.space_after = Pt(34)
contents_title = document.add_paragraph("Contenido", style="Heading 1")
contents_title.paragraph_format.space_before = Pt(18)
for line in SOURCE.read_text(encoding="utf-8").splitlines():
    match = re.match(r"^##\s+(.+)$", line)
    if not match:
        continue
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.space_after = Pt(4)
    set_font(paragraph.add_run(match.group(1)), size=10.5, color=NAVY)

body_section = document.add_section(WD_SECTION.NEW_PAGE)
body_section.page_width = Inches(8.5)
body_section.page_height = Inches(11)
body_section.top_margin = Inches(1)
body_section.right_margin = Inches(1)
body_section.bottom_margin = Inches(1)
body_section.left_margin = Inches(1)
body_section.header_distance = Inches(0.492)
body_section.footer_distance = Inches(0.492)
set_page_furniture(body_section, unlink=True)

bullet_num_id, ordered_base_num_id = set_cell_free_numbering(document)
ordered_num_id = None
list_kind = None
paragraph_lines = []


def flush_paragraph():
    if not paragraph_lines:
        return
    paragraph = document.add_paragraph()
    add_inline(paragraph, " ".join(paragraph_lines))
    paragraph_lines.clear()


lines = SOURCE.read_text(encoding="utf-8").splitlines()
for index, raw_line in enumerate(lines):
    line = raw_line.strip()
    if index < 5:
        continue
    if not line:
        flush_paragraph()
        list_kind = None
        ordered_num_id = None
        continue

    heading = re.match(r"^(#{2,3})\s+(.+)$", line)
    if heading:
        flush_paragraph()
        level = len(heading.group(1)) - 1
        document.add_paragraph(heading.group(2), style=f"Heading {level}")
        list_kind = None
        ordered_num_id = None
        continue

    unordered = re.match(r"^-\s+(.+)$", line)
    ordered = re.match(r"^\d+\.\s+(.+)$", line)
    if unordered:
        flush_paragraph()
        paragraph = document.add_paragraph()
        apply_numbering(paragraph, bullet_num_id)
        add_inline(paragraph, unordered.group(1))
        list_kind = "bullet"
        ordered_num_id = None
        continue
    if ordered:
        flush_paragraph()
        if list_kind != "ordered" or ordered_num_id is None:
            ordered_num_id = new_number_instance(document, ordered_base_num_id)
        paragraph = document.add_paragraph()
        apply_numbering(paragraph, ordered_num_id)
        add_inline(paragraph, ordered.group(1))
        list_kind = "ordered"
        continue

    list_kind = None
    ordered_num_id = None
    paragraph_lines.append(line.rstrip("  "))

flush_paragraph()
document.save(OUTPUT)
